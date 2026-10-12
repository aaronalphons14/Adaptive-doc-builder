"""Multi-format Document Parsers: Parsing raw inputs into canonical Composite trees.

Supports:
    - Structured JSON documents
    - Markdown files with YAML frontmatter
    - Plain text with administrative heading heuristics
    - DOCX binary documents via python-docx
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import yaml

from autodoc.core.models import DocumentMetadata, EntityMatch
from autodoc.core.types import DocumentType
from autodoc.patterns.composite import (
    CalloutNode,
    DocumentComponent,
    DocumentNode,
    HeadingNode,
    ListItemNode,
    ListNode,
    ParagraphNode,
    SectionNode,
    SignatureNode,
    TableCellNode,
    TableNode,
    TableRowNode,
)


class BaseParser:
    """Base parser interface."""
    def parse(self, content: str, **kwargs: Any) -> DocumentNode:
        raise NotImplementedError


class JsonDocumentParser(BaseParser):
    """Parses structured JSON representations of government documents."""

    def parse(self, content: str, **kwargs: Any) -> DocumentNode:
        data: Dict[str, Any] = json.loads(content)
        meta_dict = data.get("metadata", {})
        doc_metadata = DocumentMetadata(**meta_dict)

        doc_node = DocumentNode(doc_metadata=doc_metadata)

        for block_data in data.get("blocks", []):
            component = self._parse_block(block_data)
            if component:
                doc_node.add(component)

        return doc_node

    def _parse_block(self, b: Dict[str, Any]) -> Optional[DocumentComponent]:
        b_type = b.get("type", "").lower()

        if b_type == "section":
            sec = SectionNode(
                title=b.get("title", ""),
                section_number=b.get("number"),
            )
            for child in b.get("children", []):
                child_comp = self._parse_block(child)
                if child_comp:
                    sec.add(child_comp)
            return sec

        elif b_type == "heading":
            return HeadingNode(
                text=b.get("text", ""),
                level=b.get("level", 1),
            )

        elif b_type == "paragraph":
            entities = [EntityMatch(**e) for e in b.get("entities", [])]
            return ParagraphNode(
                text=b.get("text", ""),
                entities=entities,
                is_bilingual=b.get("is_bilingual", False),
                secondary_text=b.get("secondary_text"),
            )

        elif b_type == "table":
            headers = b.get("headers", [])
            table = TableNode(caption=b.get("caption"), headers=headers)
            if headers:
                h_row = TableRowNode(is_header_row=True)
                for h in headers:
                    h_row.add(TableCellNode(text=str(h), is_header=True))
                table.add(h_row)

            for row_data in b.get("rows", []):
                r_node = TableRowNode()
                for cell in row_data:
                    if isinstance(cell, dict):
                        r_node.add(TableCellNode(
                            text=str(cell.get("text", "")),
                            is_header=cell.get("is_header", False),
                            col_span=cell.get("col_span", 1),
                            row_span=cell.get("row_span", 1),
                        ))
                    else:
                        r_node.add(TableCellNode(text=str(cell)))
                table.add(r_node)
            return table

        elif b_type == "list":
            ordered = b.get("ordered", False)
            lst = ListNode(ordered=ordered)
            for item in b.get("items", []):
                if isinstance(item, dict):
                    lst.add(ListItemNode(text=item.get("text", ""), bullet=item.get("bullet")))
                else:
                    lst.add(ListItemNode(text=str(item)))
            return lst

        elif b_type == "signature_block" or b_type == "signature":
            return SignatureNode(
                signatory_name=b.get("name", "Authorized Signatory"),
                designation=b.get("designation", "Officer on Special Duty"),
                department=b.get("department", "Government Administration"),
                date=b.get("date"),
                digital_signed=b.get("digital_signed", False),
                sign_token=b.get("sign_token"),
            )

        elif b_type == "callout":
            return CalloutNode(
                text=b.get("text", ""),
                callout_type=b.get("callout_type", "advisory"),
                citation=b.get("citation"),
            )

        return None


class MarkdownDocumentParser(BaseParser):
    """Parses Markdown documents with YAML frontmatter."""

    def parse(self, content: str, **kwargs: Any) -> DocumentNode:
        meta_dict: Dict[str, Any] = {}
        body = content

        # Extract YAML frontmatter if present
        if content.startswith("---"):
            parts = content.split("---", 2)
            if len(parts) >= 3:
                try:
                    meta_dict = yaml.safe_load(parts[1]) or {}
                    body = parts[2]
                except Exception:
                    body = content

        doc_metadata = DocumentMetadata(**meta_dict)
        doc_node = DocumentNode(doc_metadata=doc_metadata)

        lines = body.splitlines()
        i = 0
        current_section: Optional[SectionNode] = None

        while i < len(lines):
            line = lines[i].strip()

            if not line:
                i += 1
                continue

            # Heading
            if line.startswith("#"):
                match = re.match(r"^(#+)\s*(.*)$", line)
                if match:
                    hashes, title = match.groups()
                    level = len(hashes)
                    if level == 1:
                        # Top-level heading updates document title if not set
                        if doc_node.doc_metadata.title == "Government Document":
                            doc_node.doc_metadata.title = title
                        h_node = HeadingNode(text=title, level=1)
                        doc_node.add(h_node)
                    elif level == 2:
                        current_section = SectionNode(title=title)
                        doc_node.add(current_section)
                    else:
                        h_node = HeadingNode(text=title, level=level)
                        target = current_section if current_section else doc_node
                        target.add(h_node)
                i += 1
                continue

            # Markdown Table detection
            if line.startswith("|") and line.endswith("|"):
                table_lines: List[str] = []
                while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                    table_lines.append(lines[i].strip())
                    i += 1
                table_node = self._parse_markdown_table(table_lines)
                target = current_section if current_section else doc_node
                target.add(table_node)
                continue

            # Callout / Blockquote
            if line.startswith(">"):
                callout_text = []
                while i < len(lines) and lines[i].strip().startswith(">"):
                    callout_text.append(re.sub(r"^>\s*", "", lines[i].strip()))
                    i += 1
                full_callout = " ".join(callout_text)
                citation = None
                cit_match = re.search(r"\[(Section|RTI|Order)\s*[^\]]+\]", full_callout)
                if cit_match:
                    citation = cit_match.group(0)
                callout_node = CalloutNode(text=full_callout, callout_type="statutory" if citation else "advisory", citation=citation)
                target = current_section if current_section else doc_node
                target.add(callout_node)
                continue

            # List
            if re.match(r"^(\*|-|\d+\.)\s+", line):
                list_items: List[str] = []
                is_ordered = bool(re.match(r"^\d+\.", line))
                while i < len(lines) and re.match(r"^(\*|-|\d+\.)\s+", lines[i].strip()):
                    item_text = re.sub(r"^(\*|-|\d+\.)\s+", "", lines[i].strip())
                    list_items.append(item_text)
                    i += 1
                list_node = ListNode(ordered=is_ordered)
                for item in list_items:
                    list_node.add(ListItemNode(text=item))
                target = current_section if current_section else doc_node
                target.add(list_node)
                continue

            # Signature Block heuristic in Markdown (e.g., "[Sd/- ...]")
            if "sd/-" in line.lower() or "signature:" in line.lower() or "digitally signed" in line.lower():
                sig_lines = [line]
                i += 1
                while i < len(lines) and lines[i].strip():
                    sig_lines.append(lines[i].strip())
                    i += 1
                sig_node = self._parse_signature_lines(sig_lines)
                target = current_section if current_section else doc_node
                target.add(sig_node)
                continue

            # Regular Paragraph
            target = current_section if current_section else doc_node
            target.add(ParagraphNode(text=line))
            i += 1

        return doc_node

    def _parse_markdown_table(self, lines: List[str]) -> TableNode:
        rows_data: List[List[str]] = []
        for line in lines:
            # Skip delimiter line |---|---|
            if re.match(r"^\|(?:\s*:?-+:?\s*\|)+$", line):
                continue
            cells = [c.strip() for c in line.strip("|").split("|")]
            rows_data.append(cells)

        headers = rows_data[0] if rows_data else []
        table = TableNode(headers=headers)
        if headers:
            h_row = TableRowNode(is_header_row=True)
            for h in headers:
                h_row.add(TableCellNode(text=h, is_header=True))
            table.add(h_row)

        for row_cells in rows_data[1:]:
            r_node = TableRowNode()
            for cell_text in row_cells:
                r_node.add(TableCellNode(text=cell_text))
            table.add(r_node)

        return table

    def _parse_signature_lines(self, lines: List[str]) -> SignatureNode:
        name = "Authorized Signatory"
        designation = "Competent Authority"
        department = "Government Department"
        digital_signed = any("digital" in line_str.lower() for line_str in lines)
        sign_token = None

        for line_str in lines:
            if "name:" in line_str.lower():
                name = re.sub(r"(?i)name:\s*", "", line_str)
            elif "designation:" in line_str.lower():
                designation = re.sub(r"(?i)designation:\s*", "", line_str)
            elif "department:" in line_str.lower() or "ministry:" in line_str.lower():
                department = re.sub(r"(?i)(department|ministry):\s*", "", line_str)
            elif "token:" in line_str.lower():
                sign_token = re.sub(r"(?i)token:\s*", "", line_str)
            elif not name or name == "Authorized Signatory":
                cleaned = re.sub(r"\[sd/-\]?\s*", "", line_str, flags=re.I).strip("()")
                if cleaned and not any(k in cleaned.lower() for k in ["signed", "date", "gov"]):
                    name = cleaned


        return SignatureNode(
            signatory_name=name,
            designation=designation,
            department=department,
            digital_signed=digital_signed,
            sign_token=sign_token,
        )


class PlainTextDocumentParser(BaseParser):
    """Parses plain text government documents with administrative heading heuristics."""

    def parse(self, content: str, **kwargs: Any) -> DocumentNode:
        lines = content.splitlines()
        meta = DocumentMetadata()

        # Extract metadata heuristics from header lines
        filtered_lines: List[str] = []
        for line in lines:
            stripped = line.strip()
            if not stripped:
                filtered_lines.append("")
                continue

            up = stripped.upper()
            if re.search(r"\b(GOVERNMENT ORDER|G\.O\.|ORDER)\b", up):
                meta.doc_type = DocumentType.GOVERNMENT_ORDER
                meta.title = stripped
            elif re.search(r"\b(NOTICE INVITING TENDER|NIT|TENDER)\b", up):
                meta.doc_type = DocumentType.TENDER_NOTICE
                meta.title = stripped
            elif re.search(r"\b(CERTIFICATE|BONAFIDE|DOMICILE CERTIFICATE)\b", up):
                meta.doc_type = DocumentType.CERTIFICATE
                meta.title = stripped
            elif re.search(r"\b(CIRCULAR|ADVISORY|PUBLIC NOTICE)\b", up):
                meta.doc_type = DocumentType.PUBLIC_CIRCULAR
                meta.title = stripped
            elif re.search(r"\b(RTI|RIGHT TO INFORMATION)\b", up):
                meta.doc_type = DocumentType.RTI_REPLY
                meta.title = stripped


            if re.match(r"(?i)^(No\.|Ref\s*No\.?|G\.O\.?\s*Ms\.?\s*No\.?)\s*:\s*(.*)", stripped):
                m = re.match(r"(?i)^(No\.|Ref\s*No\.?|G\.O\.?\s*Ms\.?\s*No\.?)\s*:\s*(.*)", stripped)
                if m:
                    meta.document_number = m.group(2).strip()

            filtered_lines.append(stripped)

        doc_node = DocumentNode(doc_metadata=meta)
        current_section: Optional[SectionNode] = None

        i = 0
        while i < len(filtered_lines):
            line = filtered_lines[i]
            if not line:
                i += 1
                continue

            # Check for uppercase heading heuristic (e.g. "ORDER", "SANCTION OF EXPENDITURE")
            if (
                line.isupper()
                and len(line.split()) <= 10
                and not line.startswith("NO.")
                and not line.startswith("REF")
            ):
                current_section = SectionNode(title=line)
                doc_node.add(current_section)
                i += 1
                continue

            # Numbered paragraph or section: "1. The Governor of..."
            num_match = re.match(r"^(\d+[\.\)])\s*(.*)", line)
            if num_match:
                prefix, text = num_match.groups()
                p_node = ParagraphNode(text=f"{prefix} {text}")
                target = current_section if current_section else doc_node
                target.add(p_node)
                i += 1
                continue

            # Signature block heuristic
            if any(k in line.lower() for k in ["sd/-", "signature", "competent authority", "(director)"]):
                sig_lines = [line]
                i += 1
                while i < len(filtered_lines) and filtered_lines[i]:
                    sig_lines.append(filtered_lines[i])
                    i += 1
                sig_node = MarkdownDocumentParser()._parse_signature_lines(sig_lines)
                target = current_section if current_section else doc_node
                target.add(sig_node)
                continue

            # Default paragraph
            target = current_section if current_section else doc_node
            target.add(ParagraphNode(text=line))
            i += 1

        return doc_node


class DocxDocumentParser(BaseParser):
    """Parses DOCX binary files via python-docx into the Composite tree."""

    def parse(self, content: str, **kwargs: Any) -> DocumentNode:
        try:
            import docx  # type: ignore
        except ImportError:
            # Fallback to plain text parser if python-docx unavailable
            return PlainTextDocumentParser().parse(content)

        path = Path(content)
        if not path.exists():
            # If passed raw text instead of path
            return PlainTextDocumentParser().parse(content)

        doc = docx.Document(str(path))
        meta = DocumentMetadata()
        doc_node = DocumentNode(doc_metadata=meta)
        current_section: Optional[SectionNode] = None

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style and getattr(p.style, "name", None) else ""
            if "heading 1" in style_name:

                doc_node.doc_metadata.title = text
                doc_node.add(HeadingNode(text=text, level=1))
            elif "heading 2" in style_name:
                current_section = SectionNode(title=text)
                doc_node.add(current_section)
            elif "heading" in style_name:
                target = current_section if current_section else doc_node
                target.add(HeadingNode(text=text, level=2))
            else:
                target = current_section if current_section else doc_node
                target.add(ParagraphNode(text=text))

        for table in doc.tables:
            table_node = TableNode()
            for r_idx, row in enumerate(table.rows):
                is_header = (r_idx == 0)
                row_node = TableRowNode(is_header_row=is_header)
                for cell in row.cells:
                    row_node.add(TableCellNode(text=cell.text.strip(), is_header=is_header))
                table_node.add(row_node)
            target = current_section if current_section else doc_node
            target.add(table_node)

        return doc_node


def parse_document(
    input_data: Union[str, Path],
    format_hint: Optional[str] = None,
) -> DocumentNode:
    """Unified entry point to parse any supported document into a Composite tree."""
    content: str
    path: Optional[Path] = None

    if isinstance(input_data, Path) or (isinstance(input_data, str) and Path(input_data).exists()):
        path = Path(input_data)
        if path.suffix.lower() == ".docx":
            return DocxDocumentParser().parse(str(path))
        content = path.read_text(encoding="utf-8")
        if not format_hint:
            if path.suffix.lower() == ".json":
                format_hint = "json"
            elif path.suffix.lower() in [".md", ".markdown"]:
                format_hint = "markdown"
            else:
                format_hint = "text"
    else:
        content = str(input_data)

    if not format_hint:
        trimmed = content.strip()
        if (trimmed.startswith("{") and trimmed.endswith("}")) or (trimmed.startswith("[") and trimmed.endswith("]")):
            format_hint = "json"
        elif trimmed.startswith("---") or ("\n#" in trimmed) or ("| " in trimmed):
            format_hint = "markdown"
        else:
            format_hint = "text"

    if format_hint == "json":
        return JsonDocumentParser().parse(content)
    elif format_hint == "markdown":
        return MarkdownDocumentParser().parse(content)
    elif format_hint == "docx" and path:
        return DocxDocumentParser().parse(str(path))
    else:
        return PlainTextDocumentParser().parse(content)
